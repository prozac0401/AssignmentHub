"""Set the native XLSX ID column default to text without materializing 1M cells.

Artifact Tool authors the workbook. This metadata-only export finalizer adds the
OpenXML column style, then independently reads the saved workbook to verify it.
"""
from copy import deepcopy
from pathlib import Path
import posixpath
import sys
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook


def finalize(path: Path):
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    with ZipFile(path) as archive:
        entries = [(item, archive.read(item.filename)) for item in archive.infolist()]
    content = {item.filename: data for item, data in entries}
    book = ET.fromstring(content["xl/workbook.xml"])
    first_sheet = book.find(ns + "sheets")[0]
    relation_id = first_sheet.attrib[rel_ns + "id"]
    relations = ET.fromstring(content["xl/_rels/workbook.xml.rels"])
    target = next(rel.attrib["Target"] for rel in relations if rel.attrib["Id"] == relation_id)
    sheet_path = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
    styles = ET.fromstring(content["xl/styles.xml"])
    cell_xfs = styles.find(ns + "cellXfs")
    style_index = len(cell_xfs)
    ET.SubElement(cell_xfs, ns + "xf", {"numFmtId": "49", "fontId": "0", "fillId": "0", "borderId": "0", "xfId": "0", "applyNumberFormat": "1"})
    cell_xfs.set("count", str(len(cell_xfs)))
    sheet = ET.fromstring(content[sheet_path])
    columns = sheet.find(ns + "cols")
    if columns is None:
        columns = ET.Element(ns + "cols")
        sheet.insert(list(sheet).index(sheet.find(ns + "sheetData")), columns)
    first_column = next((col for col in columns if int(col.attrib["min"]) <= 1 <= int(col.attrib["max"])), None)
    if first_column is None:
        first_column = ET.SubElement(columns, ns + "col", {"min": "1", "max": "1"})
    elif int(first_column.attrib["max"]) > 1:
        remaining = deepcopy(first_column)
        remaining.set("min", "2")
        columns.insert(list(columns).index(first_column) + 1, remaining)
        first_column.set("max", "1")
    first_column.set("style", str(style_index))
    # Artifact Tool can omit the optional dimension element. openpyxl's
    # read-only importer then reports max_row/max_column as None. Declare the
    # actual authored range so the shipped template works without an Excel
    # open/save cycle. This changes package metadata only, never cell content.
    source_book = load_workbook(path)
    dimension = sheet.find(ns + "dimension")
    if dimension is None:
        dimension = ET.Element(ns + "dimension")
        insert_at = 1 if sheet.find(ns + "sheetPr") is not None else 0
        sheet.insert(insert_at, dimension)
    dimension.set("ref", source_book.worksheets[0].calculate_dimension())
    source_book.close()
    content["xl/styles.xml"] = ET.tostring(styles, encoding="utf-8", xml_declaration=True)
    content[sheet_path] = ET.tostring(sheet, encoding="utf-8", xml_declaration=True)
    temporary = path.with_name(path.name + ".tmp")
    with ZipFile(temporary, "w") as archive:
        for item, _ in entries:
            archive.writestr(item, content[item.filename])
    temporary.replace(path)
    workbook = load_workbook(path)
    users = workbook.worksheets[0]
    assert users.column_dimensions["A"].number_format == "@"
    assert users["A2"].value == "001" and users["A2"].data_type == "s"
    assert users["A3"].value == "002" and users["A3"].data_type == "s"
    assert users["A500"].number_format == "@"
    assert users.freeze_panes == "A2"
    assert workbook.sheetnames == ["users", "작성안내"]
    workbook.close()
    streamed = load_workbook(path, read_only=True)
    assert 3 <= streamed.worksheets[0].max_row <= 10001
    assert streamed.worksheets[0].max_column == 3
    streamed.close()
    print("Verified full user_id column text default, leading zero examples, header freeze, and both sheets.")


if __name__ == "__main__":
    finalize(Path(sys.argv[1] if len(sys.argv) > 1 else "templates/users_template.xlsx"))
